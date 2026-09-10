"""
标书撰写与投标文件格式提取 API 路由模块。

暴露 API Endpoint:
- GET /documents-list                                       — 获取可用的招标文件列表
- GET /fill-bid-format/{document_id}/worker-logs            — 查取 Worker Agent 履历
- GET /fill-bid-format/{document_id}/audit-report           — Agent 填报审计报告
- GET /agent-fill-bid-format/{document_id}/download        — 下载 ReAct Agent 填报结果
- GET|POST /extract-bid-format/{document_id}              — 兼容旧调用：重新提取并下载原始模板
- POST /reextract-bid-format/{document_id}                — 重新提取并刷新原始模板缓存
- GET /download-bid-format-template/{document_id}         — 下载缓存模板，无缓存时先提取
- GET|POST /fill-bid-format/{document_id}                 — 纯净导出（不自动填报）
- POST /agent-fill-bid-format/{document_id}                — BidFillerAgent (ReAct) 填报
"""

import os
import io
import json
import re
import time
import tempfile
import uuid
import urllib.parse
from typing import Optional, Dict, Any, List
from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile, File
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import StreamingResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session
from loguru import logger

from app.api import deps
from app.db.models.user import User
from app.schemas.bid_filler_schema import BidFillRequest, RegenerateChapterRequest, RegenerateChapterResponse
from app.schemas.bid_template import (
    BidTemplateBindingRequest,
    BidTemplateBindingResponse,
    BidTemplateResponse,
)
from app.schemas.response.common import ResponseModel, success_response
from app.services.bid_format_extractor_service import bid_format_extractor_service
from app.services.bid_template_service import bid_template_service
from app.services.cost_service import resolve_cost_total
from app.services.llm_service import ModelUnavailableError

router = APIRouter()

FIRST_BID_FILL_DURATION_KEY = "first_bid_fill_duration_ms"


@router.post("/templates/upload", response_model=ResponseModel[BidTemplateResponse])
def upload_bid_template(
    file: UploadFile = File(...),
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """上传租户自己的空白 DOCX 模板，供后续绑定招标文档使用。"""
    try:
        template = bid_template_service.upload_template(
            db=db,
            file=file,
            tenant_id=current_user.tenant_id,
            user_id=current_user.id,
        )
        return success_response(data=template, message="模板上传成功")
    except ValueError as validation_error:
        logger.warning("外部投标模板上传校验失败: filename={}, error={}", file.filename, validation_error)
        raise HTTPException(status_code=400, detail=str(validation_error)) from validation_error
    except Exception as upload_error:
        logger.exception("外部投标模板上传失败: filename={}, error={}", file.filename, upload_error)
        raise HTTPException(status_code=500, detail="模板上传失败") from upload_error


@router.get("/templates", response_model=ResponseModel[List[BidTemplateResponse]])
def list_bid_templates(
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """获取当前租户可用的外部 DOCX 模板。"""
    try:
        templates = bid_template_service.list_templates(db=db, tenant_id=current_user.tenant_id)
        return success_response(data=templates)
    except Exception as list_error:
        logger.exception("查询外部投标模板失败: tenant_id={}, error={}", current_user.tenant_id, list_error)
        raise HTTPException(status_code=500, detail="查询模板列表失败") from list_error


@router.post(
    "/template-bindings/{document_id}",
    response_model=ResponseModel[BidTemplateBindingResponse],
)
def bind_bid_template(
    document_id: str,
    request_body: BidTemplateBindingRequest,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """将已上传的外部空白模板绑定到一份招标文档。"""
    try:
        binding = bid_template_service.bind_template(
            db=db,
            document_id=document_id,
            template_id=request_body.template_id,
            tenant_id=current_user.tenant_id,
            user_id=current_user.id,
            note=request_body.note,
        )
        return success_response(data=binding, message="模板绑定成功")
    except LookupError as bind_lookup_error:
        logger.warning(
            "外部投标模板绑定对象不存在或无权访问: document_id={}, template_id={}, error={}",
            document_id,
            request_body.template_id,
            bind_lookup_error,
        )
        raise HTTPException(status_code=404, detail=str(bind_lookup_error)) from bind_lookup_error
    except FileNotFoundError as template_file_error:
        logger.warning("外部投标模板物理文件缺失: template_id={}, error={}", request_body.template_id, template_file_error)
        raise HTTPException(status_code=409, detail=str(template_file_error)) from template_file_error
    except ValueError as bind_error:
        logger.warning("外部投标模板绑定失败: document_id={}, error={}", document_id, bind_error)
        raise HTTPException(status_code=400, detail=str(bind_error)) from bind_error
    except Exception as unexpected_error:
        logger.exception(
            "外部投标模板绑定发生未预期异常: document_id={}, template_id={}, error={}",
            document_id,
            request_body.template_id,
            unexpected_error,
        )
        raise HTTPException(status_code=500, detail="模板绑定失败") from unexpected_error


@router.delete(
    "/template-bindings/{document_id}",
    response_model=ResponseModel[Optional[BidTemplateBindingResponse]],
)
def unbind_bid_template(
    document_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """解除当前招标文档的外部模板绑定，后续恢复使用招标文件原格式。"""
    try:
        binding = bid_template_service.unbind_template(
            db=db,
            document_id=document_id,
            tenant_id=current_user.tenant_id,
            user_id=current_user.id,
        )
        message = "模板解除绑定成功，后续将使用招标文件原格式" if binding else "当前文档未绑定模板"
        return success_response(data=binding, message=message)
    except LookupError as unbind_lookup_error:
        logger.warning(
            "解除外部投标模板绑定对象不存在或无权访问: document_id={}, error={}",
            document_id,
            unbind_lookup_error,
        )
        raise HTTPException(status_code=404, detail=str(unbind_lookup_error)) from unbind_lookup_error
    except ValueError as unbind_error:
        logger.warning("解除外部投标模板绑定失败: document_id={}, error={}", document_id, unbind_error)
        raise HTTPException(status_code=400, detail=str(unbind_error)) from unbind_error
    except Exception as unexpected_error:
        logger.exception(
            "解除外部投标模板绑定发生未预期异常: document_id={}, error={}",
            document_id,
            unexpected_error,
        )
        raise HTTPException(status_code=500, detail="解除模板绑定失败") from unexpected_error


@router.get(
    "/template-bindings/{document_id}",
    response_model=ResponseModel[Optional[BidTemplateBindingResponse]],
)
def get_bid_template_binding(
    document_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """查询当前招标文档的外部模板绑定，未绑定时返回空数据。"""
    try:
        binding = bid_template_service.get_binding(
            db=db,
            document_id=document_id,
            tenant_id=current_user.tenant_id,
        )
        return success_response(data=binding)
    except Exception as binding_query_error:
        logger.exception(
            "查询外部投标模板绑定失败: document_id={}, tenant_id={}, error={}",
            document_id,
            current_user.tenant_id,
            binding_query_error,
        )
        raise HTTPException(status_code=500, detail="查询模板绑定失败") from binding_query_error


def _get_first_bid_fill_duration_ms(db: Session, document_id: str) -> int:
    """读取文档首次全量撰写完成时持久化的端到端耗时。"""
    from app.db.models.project import Document

    document = db.query(Document).filter(Document.id == document_id).first()
    metadata = getattr(document, "parsed_metadata", None) or {}
    duration_ms = metadata.get(FIRST_BID_FILL_DURATION_KEY) if isinstance(metadata, dict) else None
    if isinstance(duration_ms, (int, float)) and not isinstance(duration_ms, bool) and duration_ms > 0:
        return int(duration_ms)
    return 0


def _extract_manual_chapters_from_logs(logs: list) -> list[dict]:
    """从最新审计日志中读取人工撰写章节，供页面和 SSE 复用。"""
    for log in logs:
        outputs = getattr(log, "outputs", None) or {}
        manual_chapters = outputs.get("manual_chapters") if isinstance(outputs, dict) else None
        if isinstance(manual_chapters, list):
            return [item for item in manual_chapters if isinstance(item, dict)]
    return []


def _restore_profile_slots_after_chapter_reset(
    docx_path: str,
    profile: Any,
    timeline: Any,
    chapter_title: str,
) -> int:
    """在章节还原模板后，按现有字段映射补全可确认的企业档案槽位。"""
    if not docx_path or not os.path.exists(docx_path):
        return 0

    try:
        from docx import Document
        from app.agents.bid_filler_agent import _auto_fill_profile_slots
        from app.utils.table_utils import get_chapter_body_elements

        document = Document(docx_path)
        chapter_elements = get_chapter_body_elements(document, chapter_title)
        if not chapter_elements:
            logger.warning(
                "单章节重置后的档案槽位回填未找到目标章节范围: chapter={}",
                chapter_title,
            )
            return 0

        filled_count = _auto_fill_profile_slots(
            document,
            [profile] if profile is not None else [],
            timeline_source=timeline,
            allowed_elements=chapter_elements,
        )
        if filled_count:
            document.save(docx_path)
            logger.info(
                "🔄 [单章节模板回填] 章节重置后按现有字段映射补全 {} 个可确认槽位",
                filled_count,
            )
        return filled_count
    except Exception as restore_error:
        logger.exception(
            "单章节重置后的企业档案槽位回填异常: {}",
            restore_error,
        )
        return 0


def _verify_pricing_table_writeback(
    docx_path: str,
    chapter_title: str,
    proposals: List[Dict[str, Any]],
) -> tuple[bool, str]:
    """回读报价表，确认完整矩阵已经扩写到 Word 数据区。"""
    matrix_proposals = []
    for proposal in proposals or []:
        if str(proposal.get("type", "")).strip() != "table_rows":
            continue
        raw_value = proposal.get("proposed_text")
        if raw_value is None:
            raw_value = proposal.get("value", "")
        try:
            matrix = json.loads(raw_value) if isinstance(raw_value, str) else raw_value
        except json.JSONDecodeError:
            continue
        if isinstance(matrix, list) and matrix and all(isinstance(row, list) for row in matrix):
            matrix_proposals.append((proposal, matrix))

    if not matrix_proposals:
        return False, "未找到有效的报价表二维矩阵提案"

    from docx import Document
    from app.utils.table_utils import detect_table_header_rows, get_chapter_specific_table_indices

    if not docx_path or not os.path.exists(docx_path):
        return False, "报价表回读文件不存在"

    document = Document(docx_path)
    table_indices = get_chapter_specific_table_indices(document, chapter_title)
    if not table_indices:
        return False, "报价表回读时未找到目标章节表格"

    for proposal, matrix in matrix_proposals:
        path_match = re.search(r"^/body/tbl\[(\d+)\]$", str(proposal.get("path", "")).strip())
        table_index = int(path_match.group(1)) - 1 if path_match else table_indices[0]
        if table_index < 0 or table_index >= len(document.tables):
            return False, f"报价表回读路径越界: {proposal.get('path', '')}"

        table = document.tables[table_index]
        header_rows = detect_table_header_rows(table)
        expected_rows = len(matrix)
        if len(table.rows) < header_rows + expected_rows:
            return False, (
                f"报价表扩写行数不足: expected_data_rows={expected_rows}, "
                f"actual_rows={len(table.rows) - header_rows}"
            )

        for row_index in range(header_rows, header_rows + expected_rows):
            row_text = "".join(str(cell.text or "").strip() for cell in table.rows[row_index].cells)
            if not row_text:
                return False, f"报价表数据行为空: row={row_index + 1}"

    return True, f"已回读校验 {len(matrix_proposals)} 个矩阵提案"


def _query_first_bid_fill_duration_ms(document_id: str) -> int:
    """在线程池中查询首次撰写耗时，供 SSE 轮询复用独立数据库会话。"""
    from app.db.session import SessionLocal
    from sqlalchemy.exc import SQLAlchemyError

    session: Session = SessionLocal()
    try:
        return _get_first_bid_fill_duration_ms(session, document_id)
    except (SQLAlchemyError, AttributeError, TypeError, ValueError) as query_err:
        logger.exception(f"查询首次标书撰写耗时失败: document_id={document_id}, error={query_err}")
        return 0
    finally:
        session.close()


def _persist_first_bid_fill_duration(db: Session, document_id: str, duration_ms: int) -> None:
    """仅在首次成功完成全量撰写时保存耗时，后续生成不得覆盖该基准值。"""
    if duration_ms <= 0:
        logger.warning(f"首次标书撰写耗时无效，跳过持久化: document_id={document_id}, duration_ms={duration_ms}")
        return

    from app.db.models.project import Document
    from sqlalchemy.exc import SQLAlchemyError

    try:
        document = db.query(Document).filter(Document.id == document_id).first()
        if document is None:
            logger.warning(f"未找到文档，无法持久化首次标书撰写耗时: document_id={document_id}")
            return

        metadata = dict(getattr(document, "parsed_metadata", None) or {})
        existing_duration = metadata.get(FIRST_BID_FILL_DURATION_KEY)
        if isinstance(existing_duration, (int, float)) and not isinstance(existing_duration, bool) and existing_duration > 0:
            logger.info(f"首次标书撰写耗时已存在，保持原值: document_id={document_id}, duration_ms={int(existing_duration)}")
            return

        metadata[FIRST_BID_FILL_DURATION_KEY] = int(duration_ms)
        document.parsed_metadata = metadata
        db.commit()
        logger.info(f"已持久化首次标书撰写耗时: document_id={document_id}, duration_ms={duration_ms}")
    except (SQLAlchemyError, AttributeError, TypeError, ValueError) as persist_err:
        db.rollback()
        logger.exception(f"持久化首次标书撰写耗时失败: document_id={document_id}, error={persist_err}")


def _get_bid_fill_pipeline_state(logs: list) -> Dict[str, Any]:
    """根据最终 Supervisor 终态日志判断整条标书填报流水线状态。"""
    def _log_sort_key(log: Any) -> float:
        """统一转换日志时间，兼容数据库时间对象和测试替身。"""
        created_at = getattr(log, "created_at", None)
        if hasattr(created_at, "timestamp"):
            return float(created_at.timestamp())
        if isinstance(created_at, (int, float)):
            return float(created_at)
        return 0.0

    terminal_logs = [
        log for log in logs
        if getattr(log, "node_name", "") == "Supervisor-总控调度"
        and getattr(log, "status", "") in {"master_completed", "failed", "error"}
    ]
    latest_terminal = max(
        terminal_logs,
        key=_log_sort_key,
        default=None,
    )
    if latest_terminal is None:
        return {
            "pipeline_status": "processing" if logs else "idle",
            "pipeline_message": "后台 Agent 正在执行章节填报与终审校验",
            "is_completed": False,
        }

    status = getattr(latest_terminal, "status", "")
    if status == "master_completed":
        return {
            "pipeline_status": "completed",
            "pipeline_message": "后台填报、终审和最终 Word 发布流程已完成",
            "is_completed": True,
        }

    # 模型连接失败时，优先返回稳定的业务提示，避免前端只能看到笼统的流程异常。
    failure_messages: list[str] = []
    error_message = getattr(latest_terminal, "error_message", None)
    if isinstance(error_message, str) and error_message.strip():
        failure_messages.append(error_message.strip())
    outputs = getattr(latest_terminal, "outputs", None)
    if isinstance(outputs, dict):
        for key in ("error", "error_message", "summary"):
            value = outputs.get(key)
            if isinstance(value, str) and value.strip():
                failure_messages.append(value.strip())

    normalized_failure = " ".join(failure_messages).lower()
    if (
        "模型不可用" in normalized_failure
        or "模型服务不可用" in normalized_failure
        or "model unavailable" in normalized_failure
        or (
            "model service" in normalized_failure
            and "connection" in normalized_failure
        )
    ) and (
        "连接失败" in normalized_failure
        or "服务连接" in normalized_failure
        or "connection" in normalized_failure
        or "disconnected" in normalized_failure
    ):
        return {
            "pipeline_status": "failed",
            "pipeline_message": "模型服务不可用：模型服务连接失败，请检查模型地址、网络或服务状态",
            "is_completed": True,
        }

    return {
        "pipeline_status": "failed",
        "pipeline_message": "后台填报流程异常结束，请查看审计日志",
        "is_completed": True,
    }


def _query_bid_fill_logs(document_id: str) -> list[Any]:
    """在线程池中查询标书撰写日志，避免同步数据库 I/O 阻塞事件循环。"""
    from app.db.models.audit import AgentAuditLog
    from app.db.session import SessionLocal
    from sqlalchemy import cast, desc, String

    session: Session = SessionLocal()
    try:
        try:
            return (
                session.query(AgentAuditLog)
                .filter(
                    or_(
                        AgentAuditLog.task_id == document_id,
                        cast(AgentAuditLog.inputs, String).like(f"%{document_id}%")
                    )
                )
                .order_by(desc(AgentAuditLog.created_at))
                .all()
            )
        except Exception as filter_err:
            # 兼容历史数据库 JSON 字段类型不支持 CAST LIKE 的情况。
            logger.warning(f"基于 SQL 过滤 AgentAuditLog 异常，降级全量过滤: {filter_err}")
            all_logs = (
                session.query(AgentAuditLog)
                .order_by(desc(AgentAuditLog.created_at))
                .limit(200)
                .all()
            )
            return [
                log for log in all_logs
                if log.task_id == document_id or document_id in str(log.inputs or {})
            ]
    except Exception as query_err:
        logger.exception(f"查询标书撰写日志失败: document_id={document_id}, error={query_err}")
        return []
    finally:
        session.close()


# ============================================================
# 1. 静态及特定多层子路径 Endpoint (必须位于单层 {document_id} 之前)
# ============================================================

@router.get("/documents-list")
def get_bidding_documents_list(
    doc_type: Optional[str] = None,
    db: Session = Depends(deps.get_db),
    current_user: Optional[User] = Depends(deps.get_current_user_optional)
):
    """
    获取系统中已上传并解析的全套招标文件列表，供前端智能撰写控制台下拉框选择
    默认仅返回 doc_type=tender 的招标文件
    """
    from app.db.crud.document import document_crud
    user_id = current_user.id if current_user else None
    tenant_id = current_user.tenant_id if current_user else None
    
    # 默认针对标书撰写拉取仅招标文件
    target_type = doc_type if doc_type else "tender"
    
    docs = document_crud.get_all_documents(db, user_id, tenant_id, doc_type=target_type)
    if not docs or len(docs) == 0:
        docs = document_crud.get_all_documents(db, None, None, doc_type=target_type)

    res = []
    for d in docs:
        pm = d.parsed_metadata or {}
        proj_name = pm.get("project_name") or pm.get("title") or pm.get("name")
        proj_code = pm.get("project_code")
        
        # 尝试从关联的项目 Project 获取名称
        if not proj_name and hasattr(d, "project") and d.project:
            proj_name = d.project.name

        # 尝试从切片标题提取人类可读项目名
        if not proj_name and hasattr(d, "chunks") and d.chunks:
            for c in d.chunks[:5]:
                if c.section_title and len(c.section_title) > 2 and "目录" not in c.section_title:
                    proj_name = c.section_title
                    break

        code_str = f"[{proj_code}] " if proj_code else ""
        display_label = f"{code_str}{proj_name} ({d.filename})" if proj_name else d.filename

        res.append({
            "id": d.id,
            "filename": d.filename,
            "doc_type": pm.get("doc_type", "tender"),
            "project_name": proj_name or d.filename,
            "project_code": proj_code or "--",
            "display_label": display_label,
            "parse_status": d.parse_status or "completed",
            "created_at": d.created_at.isoformat() if hasattr(d, "created_at") and d.created_at else None
        })
    return res


@router.get("/fill-bid-format/{document_id}/worker-logs")
def get_bid_fill_worker_logs(
    document_id: str,
    db: Session = Depends(deps.get_db),
    current_user: Optional[User] = Depends(deps.get_current_user_optional)
):
    """
    直查数据库获取全套 BidFillerWorker 子 Agent 节点真实运行履历与思考推导总结
    """
    if not document_id:
        raise HTTPException(status_code=400, detail="未提供有效的 document_id 参数")

    try:
        from app.db.models.audit import AgentAuditLog
        from sqlalchemy import desc, cast, String

        try:
            logs = (
                db.query(AgentAuditLog)
                .filter(
                    or_(
                        AgentAuditLog.task_id == document_id,
                        cast(AgentAuditLog.inputs, String).like(f"%{document_id}%")
                    )
                )
                .order_by(desc(AgentAuditLog.created_at))
                .all()
            )
        except Exception as filter_err:
            logger.warning(f"基于 SQL 过滤 AgentAuditLog 异常, 降级全量过滤: {filter_err}")
            all_logs = db.query(AgentAuditLog).order_by(desc(AgentAuditLog.created_at)).limit(200).all()
            logs = [l for l in all_logs if l.task_id == document_id or document_id in str(l.inputs or {})]

        worker_items = []
        seen_chapters = set()

        first_bid_fill_duration_ms = _get_first_bid_fill_duration_ms(db, document_id)
        total_wall_time_ms = first_bid_fill_duration_ms
        min_created_at = None
        max_created_at = None

        for log in logs:
            if log.created_at:
                if min_created_at is None or log.created_at < min_created_at:
                    min_created_at = log.created_at
                if max_created_at is None or log.created_at > max_created_at:
                    max_created_at = log.created_at

            if total_wall_time_ms == 0 and log.status == "master_completed" and log.execution_time_ms and log.execution_time_ms > 0:
                total_wall_time_ms = max(total_wall_time_ms, log.execution_time_ms)

            if log.action_type in ("llm_call_worker", "llm_call_supervisor", "chapter_execution") or (log.node_name and (log.node_name.startswith("BidFillerWorker") or "Supervisor" in log.node_name)):
                inp = log.inputs or {}
                out = log.outputs or {}
                is_supervisor = (log.node_name and "Supervisor" in log.node_name) or log.action_type == "llm_call_supervisor"
                
                ch_title = "Supervisor 总控调度" if is_supervisor else (inp.get("chapter_title") or (log.node_name.replace("BidFillerWorker-", "") if log.node_name else "未知章节"))

                if ch_title in seen_chapters:
                    continue
                seen_chapters.add(ch_title)

                worker_items.append({
                    "id": str(log.id),
                    "node_name": "Supervisor-总控调度" if is_supervisor else (log.node_name or f"BidFillerWorker-{ch_title}"),
                    "chapter_title": ch_title,
                    "category": "supervisor_master" if is_supervisor else inp.get("category", "needs_fill"),
                    "status": log.status or "success",
                    "execution_time_ms": log.execution_time_ms or 0,
                    "total_tokens": log.total_tokens or ((log.prompt_tokens or 0) + (log.completion_tokens or 0)),
                    "prompt_tokens": log.prompt_tokens or 0,
                    "completion_tokens": log.completion_tokens or 0,
                    "summary": out.get("summary", "已完成填报分析与写盘。"),
                    "proposals_count": out.get("proposals_count", 0),
                    "written_count": out.get("written_count"),
                    "proposals": out.get("proposals", []),
                    "tools_used": out.get("tools_used", inp.get("tools_used", [])),
                    "thought_steps": out.get("thought_steps", []),
                    "created_at": log.created_at.strftime("%Y-%m-%d %H:%M:%S") if log.created_at else None
                })

        total_worker_time_ms = sum(w.get("execution_time_ms", 0) for w in worker_items)
        if total_wall_time_ms == 0 and min_created_at and max_created_at:
            delta_ms = int((max_created_at - min_created_at).total_seconds() * 1000)
            if delta_ms > 0:
                total_wall_time_ms = delta_ms

        pipeline_state = _get_bid_fill_pipeline_state(logs)
        manual_chapters = _extract_manual_chapters_from_logs(logs)
        return {
            "document_id": document_id,
            "total_workers_count": len(worker_items),
            "worker_items": worker_items,
            "manual_chapters": manual_chapters,
            "manual_pending_count": len(manual_chapters),
            "total_wall_time_ms": total_wall_time_ms,
            "first_bid_fill_duration_ms": first_bid_fill_duration_ms,
            "total_worker_time_ms": total_worker_time_ms,
            **pipeline_state,
        }
    except Exception as e:
        logger.exception(f"获取 Agent 运行日志出现异常: {e}")
        return {
            "document_id": document_id,
            "total_workers_count": 0,
            "worker_items": [],
            "manual_chapters": [],
            "manual_pending_count": 0,
            "total_wall_time_ms": 0,
            "first_bid_fill_duration_ms": 0,
            "total_worker_time_ms": 0
        }


def _build_bid_fill_stream_payload(document_id: str) -> tuple[str, bool]:
    """在线程池内查询、组装并序列化 SSE 快照，避免大日志占用事件循环。"""
    logs = _query_bid_fill_logs(document_id)
    worker_items = []
    seen_chapters = set()

    first_bid_fill_duration_ms = _query_first_bid_fill_duration_ms(document_id)
    total_wall_time_ms = first_bid_fill_duration_ms
    min_created_at = None
    max_created_at = None

    for log in logs:
        if log.created_at:
            if min_created_at is None or log.created_at < min_created_at:
                min_created_at = log.created_at
            if max_created_at is None or log.created_at > max_created_at:
                max_created_at = log.created_at

        if total_wall_time_ms == 0 and log.status == "master_completed" and log.execution_time_ms and log.execution_time_ms > 0:
            total_wall_time_ms = max(total_wall_time_ms, log.execution_time_ms)

        if log.action_type in ("llm_call_worker", "llm_call_supervisor", "chapter_execution") or (log.node_name and (log.node_name.startswith("BidFillerWorker") or "Supervisor" in log.node_name)):
            inp = log.inputs or {}
            out = log.outputs or {}
            is_supervisor = (log.node_name and "Supervisor" in log.node_name) or log.action_type == "llm_call_supervisor"
            ch_title = "Supervisor 总控调度" if is_supervisor else (inp.get("chapter_title") or (log.node_name.replace("BidFillerWorker-", "") if log.node_name else "未知章节"))

            if ch_title in seen_chapters:
                continue
            seen_chapters.add(ch_title)

            status_val = log.status or "success"

            worker_items.append({
                "id": str(log.id),
                "node_name": "Supervisor-总控调度" if is_supervisor else (log.node_name or f"BidFillerWorker-{ch_title}"),
                "chapter_title": ch_title,
                "category": "supervisor_master" if is_supervisor else inp.get("category", "needs_fill"),
                "status": status_val,
                "execution_time_ms": log.execution_time_ms or 0,
                "total_tokens": log.total_tokens or ((log.prompt_tokens or 0) + (log.completion_tokens or 0)),
                "prompt_tokens": log.prompt_tokens or 0,
                "completion_tokens": log.completion_tokens or 0,
                "summary": out.get("summary", "已完成填报分析与写盘。"),
                "proposals_count": out.get("proposals_count", 0),
                "written_count": out.get("written_count"),
                "proposals": out.get("proposals", []),
                "tools_used": out.get("tools_used", inp.get("tools_used", [])),
                "thought_steps": out.get("thought_steps", []),
                "created_at": log.created_at.strftime("%Y-%m-%d %H:%M:%S") if log.created_at else None,
            })

    total_worker_time_ms = sum(w.get("execution_time_ms", 0) for w in worker_items)
    if total_wall_time_ms == 0 and min_created_at and max_created_at:
        delta_ms = int((max_created_at - min_created_at).total_seconds() * 1000)
        if delta_ms > 0:
            total_wall_time_ms = delta_ms

    pipeline_state = _get_bid_fill_pipeline_state(logs)
    manual_chapters = _extract_manual_chapters_from_logs(logs)
    payload = {
        "document_id": document_id,
        "worker_items": worker_items,
        "manual_chapters": manual_chapters,
        "manual_pending_count": len(manual_chapters),
        # 只有后台最终 Supervisor 终态才能结束 SSE，不能用中间 Worker 成功状态代替。
        "is_completed": pipeline_state["is_completed"],
        "pipeline_status": pipeline_state["pipeline_status"],
        "pipeline_message": pipeline_state["pipeline_message"],
        "total_wall_time_ms": total_wall_time_ms,
        "first_bid_fill_duration_ms": first_bid_fill_duration_ms,
        "total_worker_time_ms": total_worker_time_ms,
        "timestamp": time.time(),
    }
    return json.dumps(payload, ensure_ascii=False), bool(pipeline_state["is_completed"])


@router.get("/fill-bid-format/{document_id}/stream-logs")
async def stream_bid_fill_worker_logs(document_id: str):
    """通过 SSE 实时推流获取 BidFillerWorker 全套 Agent 节点运行履历。"""
    if not document_id:
        raise HTTPException(status_code=400, detail="未提供有效的 document_id 参数")

    async def log_event_generator():
        import asyncio

        last_json = None
        same_count = 0

        while True:
            try:
                # 查询、组装和 JSON 序列化全部在线程池执行，避免大日志阻塞其他页面请求。
                payload_str, is_completed = await run_in_threadpool(
                    _build_bid_fill_stream_payload,
                    document_id,
                )

                if payload_str != last_json:
                    last_json = payload_str
                    same_count = 0
                    yield f"data: {payload_str}\n\n"
                else:
                    same_count += 1
                    yield f": ping {int(time.time())}\n\n"

                if is_completed and same_count >= 5:
                    break
            except Exception as stream_error:
                logger.error(f"SSE 推流日志生成异常: {stream_error}")

            await asyncio.sleep(1.0)

    return StreamingResponse(
        log_event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )


@router.post("/fill-bid-format/{document_id}/regenerate-chapter", response_model=RegenerateChapterResponse)
async def regenerate_single_chapter(
    document_id: str,
    request_body: RegenerateChapterRequest,
    db: Session = Depends(deps.get_db),
    current_user: Optional[User] = Depends(deps.get_current_user_optional)
):
    """
    针对具体指定章节重新启动 Worker Agent 进行针对性起草与 Prompt 微调，并原位写回 Word 文档。
    """
    if not document_id:
        raise HTTPException(status_code=400, detail="未提供有效的 document_id 参数")
    if not request_body.chapter_title:
        raise HTTPException(status_code=400, detail="未提供目标章节名称 chapter_title")

    chapter_title = request_body.chapter_title.strip()
    custom_prompt = (request_body.custom_prompt or "").strip()
    category = (request_body.category or "needs_fill").strip()
    mapping_hint = (request_body.mapping_hint or "").strip()

    logger.info(
        f"🔄 收到单章节重新生成/微调请求: doc_id={document_id}, "
        f"chapter={chapter_title}, profile_id={request_body.profile_id}, "
        f"prompt='{custom_prompt[:60]}'"
    )

    # 1. 准备 Word 工作副本与纯净原始模板
    drafts_dir = os.path.join(os.getcwd(), "uploads", "drafts")
    os.makedirs(drafts_dir, exist_ok=True)
    working_docx_path = os.path.join(drafts_dir, f"bid_fill_{document_id[:8]}.docx")
    result_docx_path = os.path.join(drafts_dir, f"agent_fill_result_{document_id[:8]}.docx")
    template_docx_path = os.path.join(drafts_dir, f"template_{document_id[:8]}.docx")

    # 确保纯净模板存在（供单章微调提取纯净无污染的占位符上下文）
    if not os.path.exists(template_docx_path):
        template_bytes, filename, _ = bid_format_extractor_service.extract_and_export_bid_format(
            db=db, doc_id=document_id, user_id=None, tenant_id=None
        )
        if template_bytes:
            with open(template_docx_path, "wb") as f:
                f.write(template_bytes)

    if not os.path.exists(working_docx_path):
        if os.path.exists(result_docx_path):
            import shutil
            shutil.copyfile(result_docx_path, working_docx_path)
        elif os.path.exists(template_docx_path):
            import shutil
            shutil.copyfile(template_docx_path, working_docx_path)

    # 2. 注入上下文变量
    from app.core.context import current_user_id, current_tenant_id, current_task_id
    u_id = current_user.id if (current_user and hasattr(current_user, 'id')) else "default-user"
    t_id = current_user.tenant_id if (current_user and hasattr(current_user, 'tenant_id')) else "default-tenant"
    token_task = current_task_id.set(document_id)
    token_u = current_user_id.set(u_id)
    token_t = current_tenant_id.set(t_id)
    # 单章节重生成与全量撰写保持一致，显式绑定本次选择的企业主体。
    from app.agents.tools.bid_db_tools import current_profile_id as ctx_profile_id
    token_profile = ctx_profile_id.set(request_body.profile_id)

    # 3. 记录初始进行中状态
    try:
        from app.services.audit_service import audit_service
        audit_service.log_event(
            action_type="llm_call_worker",
            node_name=f"BidFillerWorker-{chapter_title[:30]}",
            inputs={
                "chapter_title": chapter_title,
                "category": category,
                "document_id": document_id,
                "profile_id": request_body.profile_id,
                "custom_prompt": custom_prompt,
            },
            outputs={
                "summary": f"🔄 正在根据用户提示词对章节 [{chapter_title}] 重新起草与微调...",
                "proposals_count": 0,
                "thought_steps": [
                    {"step": 1, "type": "thought", "content": f"接收到用户微调指令: '{custom_prompt}'，正在启动专属 Worker Agent 重新调取数据并生成提案。"}
                ]
            },
            status="in_progress"
        )
    except Exception as log_init_err:
        logger.warning(f"写入微调初始状态日志异常: {log_init_err}")

    try:
        from app.agents.bid_filler_workers import _is_pricing_chapter, run_chapter_worker
        from app.agents.bid_filler_agent import fill_docx_proposals_in_dom
        from app.utils.table_utils import extract_chapter_dom_structure

        # 优先从纯净模板提取章节专属上下文，并在重新生成前将工作副本中该章节重置为纯净模板状态（确保100%覆盖）
        from app.agents.bid_filler_workers import extract_docx_tables_summary
        from app.utils.table_utils import get_chapter_specific_table_indices, reset_chapter_to_template

        if os.path.exists(template_docx_path) and os.path.exists(working_docx_path):
            reset_chapter_to_template(working_docx_path, template_docx_path, chapter_title)

        # 章节重置后，工作副本中的目标章节已为模板纯净状态，且绝对段落与表格索引
        # 与当前文档拓扑（包含前面章节插入的图片与段落）100% 一致。
        # 优先从工作副本提取章节 DOM 视野，确保 Worker 生成的物理路径与写盘坐标系绝对对齐。
        template_source_path = working_docx_path if os.path.exists(working_docx_path) else template_docx_path
        target_tbl_summary = extract_docx_tables_summary(template_source_path, chapter_title)
        
        # 组装纯净模板提示（如为表格章节，突出表头定义与全量重写要求）
        if target_tbl_summary:
            chapter_pure_context = f"【本章节专属表格表头定义】\n{target_tbl_summary}\n（请根据招标文件原文及企业数据库全量检索数据，生成完整 2D 矩阵全量覆写）"
        else:
            chapter_pure_context = extract_chapter_dom_structure(template_source_path, chapter_title)
            if not chapter_pure_context and template_source_path != template_docx_path and os.path.exists(template_docx_path):
                chapter_pure_context = extract_chapter_dom_structure(template_docx_path, chapter_title)
            if not chapter_pure_context:
                chapter_pure_context = f"【目标章节】: {chapter_title}"

        # 尝试从历史日志中继承该章节的分类 hint 与说明
        effective_category = category
        effective_mapping_hint = mapping_hint
        try:
            from app.db.models.audit import AgentAuditLog
            from sqlalchemy import desc
            hist_log = (
                db.query(AgentAuditLog)
                .filter(AgentAuditLog.task_id == document_id)
                .filter(AgentAuditLog.node_name == f"BidFillerWorker-{chapter_title[:30]}")
                .order_by(desc(AgentAuditLog.created_at))
                .first()
            )
            if hist_log and hist_log.inputs:
                if not effective_mapping_hint:
                    effective_mapping_hint = hist_log.inputs.get("mapping_hint", "")
                if not effective_category or effective_category == "needs_fill":
                    effective_category = hist_log.inputs.get("category", effective_category)
        except Exception:
            pass

        # 预读取企业档案与项目元数据（用于公文类单章微调定向注入）
        prefetched_metadata: Dict[str, Any] = {}
        prof = None
        tl = None
        try:
            from app.agents.tools.bid_db_tools import resolve_company_profile
            from app.db.models.metadata import TimelineMetadata, FinancialMetadata
            from app.utils.date_formatter import format_date_only
            from app.utils.rmb_formatter import number_to_chinese_rmb

            # 读取指定主体，禁止通过无序 first() 串用其他企业档案。
            prof = resolve_company_profile(db, request_body.profile_id)
            if prof:
                if prof.company_name: prefetched_metadata["company_name"] = prof.company_name
                if prof.credit_code: prefetched_metadata["credit_code"] = prof.credit_code
                if prof.legal_representative: prefetched_metadata["legal_person"] = prof.legal_representative
                if prof.registered_address: prefetched_metadata["address"] = prof.registered_address
                if prof.contact_phone: prefetched_metadata["phone"] = prof.contact_phone
                if prof.email: prefetched_metadata["email"] = prof.email

            tl = db.query(TimelineMetadata).filter(TimelineMetadata.document_id == document_id).first()
            if tl:
                if getattr(tl, "project_name", None): prefetched_metadata["project_name"] = tl.project_name
                proj_code = getattr(tl, "project_id_code", None) or getattr(tl, "project_code", None)
                if proj_code: prefetched_metadata["project_code"] = proj_code
                deadline_date = format_date_only(getattr(tl, "bid_deadline", None))
                if deadline_date:
                    prefetched_metadata["bid_deadline_date"] = deadline_date
                
                period_str = str(getattr(tl, "construction_period_description", "") or "").strip()
                if not period_str and getattr(tl, "construction_period_days", None):
                    period_str = f"{tl.construction_period_days} 日历天"
                if period_str:
                    prefetched_metadata["delivery_period"] = period_str

            from app.db.models.ai_analysis import CostEstimate
            from app.db.models.project import Document
            cost_items = db.query(CostEstimate).filter(CostEstimate.document_id == document_id).all()
            cost_document = db.query(Document).filter(Document.id == document_id).first()
            stored_cost_analysis = (
                (cost_document.parsed_metadata or {}).get("cost_analysis", {})
                if cost_document and isinstance(cost_document.parsed_metadata, dict)
                else {}
            )
            total_val = resolve_cost_total(stored_cost_analysis, cost_items)
            if total_val > 0:
                prefetched_metadata["total_price_str"] = f"{total_val:,.2f} 元"
                try:
                    prefetched_metadata["total_price_words"] = number_to_chinese_rmb(float(total_val))
                except (TypeError, ValueError):
                    logger.warning("微调接口成本总额转人民币大写失败：{}", total_val)

            prefetched_metadata["quality_standard"] = "合格，完全符合国家及行业现行有关标准、规范要求"
        except Exception as e_meta:
            logger.warning(f"微调接口预读取企业与项目元数据异常: {e_meta}")

        # Worker 提案数量受模型判断影响，不能用它作为模板槽位完整性的唯一保证。
        # 章节已还原为干净模板后，先用统一字段映射回填已确认的档案值，再交给 Worker
        # 处理需要语义判断的内容；这样不会依赖某一种表单名称或固定段落编号。
        restored_slot_count = _restore_profile_slots_after_chapter_reset(
            working_docx_path,
            prof,
            tl,
            chapter_title,
        )

        start_time = time.time()
        worker_res = await run_in_threadpool(
            run_chapter_worker,
            chapter_title=chapter_title,
            chapter_number="",
            mapping_hint=effective_mapping_hint,
            category=effective_category,
            document_id=document_id,
            docx_temp_path=working_docx_path,
            template_text=chapter_pure_context,
            content_hint="（请根据招标文件原文与企业数据库检索全量数据，按要求全量重新起草与覆写本章节）",
            extra_instructions=custom_prompt or "请按照主流程标准，全量重新检索招标文件与数据库并完成全表覆写。",
            repair_instructions="",
            prefetched_metadata=prefetched_metadata,
            tenant_id=t_id,
            profile_id=request_body.profile_id,
        )
        elapsed_ms = int((time.time() - start_time) * 1000)

        proposals = worker_res.get("proposals", [])
        status = worker_res.get("status", "success")
        summary = worker_res.get("summary", "单章节微调已完成。")
        is_pricing_request = _is_pricing_chapter(chapter_title, effective_mapping_hint)
        write_count = 0
        writeback_verified = not is_pricing_request

        # 4. 原位刷盘
        if os.path.exists(working_docx_path):
            try:
                write_count = fill_docx_proposals_in_dom(working_docx_path, proposals) if proposals else 0
                if is_pricing_request:
                    if status == "success" and write_count > 0:
                        writeback_verified, verify_message = _verify_pricing_table_writeback(
                            working_docx_path,
                            chapter_title,
                            proposals,
                        )
                        if not writeback_verified:
                            status = "failed"
                            summary = f"报价表写盘回读校验失败：{verify_message}"
                            logger.error(f"❌ {summary}")
                        else:
                            logger.info(f"✅ 报价表确定性闭环完成: {verify_message}")
                    else:
                        status = "failed"
                        summary = "报价 Worker 未产生可写入的有效提案，未保存章节结果。"
                        logger.error(f"❌ {summary}")

                logger.info(f"✅ 单章节微调原位写盘完成，写入 {write_count} 处修改")
                # 只有确定性闭环通过后才同步到结果文件，避免把重置后的空模板覆盖正式草稿。
                should_persist_result = (
                    status == "success"
                    and (not is_pricing_request or writeback_verified)
                )
                if should_persist_result:
                    import shutil
                    shutil.copyfile(working_docx_path, result_docx_path)
                    draft_path = os.path.join(drafts_dir, f"draft_{document_id}.docx")
                    shutil.copyfile(working_docx_path, draft_path)
                else:
                    logger.warning(
                        f"⚠️ 单章节微调未通过结果保存条件: status={status}, "
                        f"pricing={is_pricing_request}, verified={writeback_verified}"
                    )
            except Exception as write_err:
                status = "failed"
                summary = f"单章节微调写盘或回读异常：{write_err}"
                logger.error(f"单章节微调写盘异常: {write_err}")

        if is_pricing_request and status != "success":
            raise HTTPException(status_code=422, detail=summary)

        # 5. 查询最新的 audit log 条目
        from app.db.models.audit import AgentAuditLog
        from sqlalchemy import desc, cast, String
        latest_log = (
            db.query(AgentAuditLog)
            .filter(
                or_(
                    AgentAuditLog.task_id == document_id,
                    cast(AgentAuditLog.inputs, String).like(f"%{document_id}%")
                )
            )
            .filter(
                or_(
                    AgentAuditLog.node_name == f"BidFillerWorker-{chapter_title[:30]}",
                    cast(AgentAuditLog.inputs, String).like(f"%{chapter_title}%")
                )
            )
            .order_by(desc(AgentAuditLog.created_at))
            .first()
        )

        worker_item = None
        if latest_log:
            inp = latest_log.inputs or {}
            out = latest_log.outputs or {}
            # 将本次 DOM 写盘后的逐条状态回写审计日志，避免前端把“提案成功”误显示成“物理写盘成功”。
            updated_outputs = dict(out)
            updated_outputs["proposals"] = proposals
            updated_outputs["proposals_count"] = len(proposals)
            updated_outputs["written_count"] = write_count
            latest_log.outputs = updated_outputs
            try:
                db.commit()
                out = updated_outputs
            except Exception as audit_error:
                db.rollback()
                logger.exception(f"单章节写盘结果回写审计日志失败: {audit_error}")
                # 数据库回写失败不影响已生成的 Word 文件，接口仍使用本次内存结果返回。
                out = updated_outputs
            worker_item = {
                "id": str(latest_log.id),
                "node_name": latest_log.node_name or f"BidFillerWorker-{chapter_title}",
                "chapter_title": chapter_title,
                "category": inp.get("category", category),
                "status": latest_log.status or "success",
                "execution_time_ms": latest_log.execution_time_ms or elapsed_ms,
                "total_tokens": latest_log.total_tokens or 0,
                "prompt_tokens": latest_log.prompt_tokens or 0,
                "completion_tokens": latest_log.completion_tokens or 0,
                "summary": out.get("summary", summary),
                "proposals_count": out.get("proposals_count", len(proposals)),
                "written_count": out.get("written_count", write_count),
                "proposals": out.get("proposals", proposals),
                "tools_used": out.get("tools_used", inp.get("tools_used", [])),
                "thought_steps": out.get("thought_steps", []),
                "created_at": latest_log.created_at.strftime("%Y-%m-%d %H:%M:%S") if latest_log.created_at else None
            }

        return RegenerateChapterResponse(
            document_id=document_id,
            chapter_title=chapter_title,
            status=status,
            summary=summary,
            proposals_count=len(proposals),
            written_count=write_count,
            execution_time_ms=elapsed_ms,
            total_tokens=worker_item.get("total_tokens", 0) if worker_item else 0,
            worker_item=worker_item
        )

    except HTTPException:
        raise
    except Exception as exc:
        logger.exception(f"❌ 单章节微调失败: {exc}")
        raise HTTPException(status_code=500, detail=f"单章节重新生成失败: {str(exc)}")
    finally:
        try:
            current_task_id.reset(token_task)
            current_user_id.reset(token_u)
            current_tenant_id.reset(token_t)
            ctx_profile_id.reset(token_profile)
        except Exception:
            pass



@router.get("/fill-bid-format/{document_id}/audit-report")
async def get_bid_fill_audit_report(
    document_id: str,
    db: Session = Depends(deps.get_db),
    current_user: Optional[User] = Depends(deps.get_current_user_optional)
):
    """
    获取 Agent 标书填报对齐追溯核查报告 (Filling Audit Trail Report)
    包含字段级对齐报告 + 数据库保存的子 Agent 真实思考与执行履历
    """
    from app.schemas.bid_filler_schema import CompanyProfile
    from app.agents.bid_filler_agent import bid_filler_agent
    template_bytes, filename, _ = await run_in_threadpool(
        bid_format_extractor_service.extract_and_export_bid_format,
        db=db,
        doc_id=document_id,
        user_id=current_user.id if hasattr(current_user, "id") else None,
        tenant_id=current_user.tenant_id if hasattr(current_user, "tenant_id") else None,
    )

    if not template_bytes:
        raise HTTPException(status_code=404, detail="未找到该文档的模版信息")

    # Agent 自行通过 OfficeCLI 阅读 Word 文档发现需要填写的位置
    _, audit_report, _ = await run_in_threadpool(
        bid_filler_agent.process_filling_tasks,
        db=db,
        document_id=document_id,
        profile=CompanyProfile(),
        detected_placeholders=[],
        original_docx=template_bytes,
    )

    res_dict = audit_report.model_dump() if hasattr(audit_report, 'model_dump') else audit_report.dict()

    # 融合直查数据库得到的子 Agent 思考全过程履历
    try:
        worker_logs = await run_in_threadpool(
            get_bid_fill_worker_logs,
            document_id=document_id,
            db=db,
            current_user=current_user,
        )
        res_dict["worker_items"] = worker_logs.get("worker_items", [])
        res_dict["total_workers_count"] = worker_logs.get("total_workers_count", 0)
        res_dict["manual_chapters"] = worker_logs.get("manual_chapters", [])
        res_dict["manual_pending_count"] = worker_logs.get("manual_pending_count", 0)
    except Exception as exc:
        logger.warning(f"获取 Worker 审计日志融合失败: {exc}")
        res_dict["worker_items"] = []
        res_dict["total_workers_count"] = 0
        res_dict["manual_chapters"] = []
        res_dict["manual_pending_count"] = 0

    return res_dict


def _has_non_empty_result_file(result_path: str) -> bool:
    """判断标书撰写结果文件是否已经生成且不是空文件。"""
    try:
        return os.path.isfile(result_path) and os.path.getsize(result_path) > 0
    except OSError as file_error:
        logger.warning(f"检查标书撰写结果文件失败: path={result_path}, error={file_error}")
        return False


def _read_agent_result_file(
    result_path: str,
    document_id: str,
    user_id: Optional[str],
    tenant_id: Optional[str],
) -> tuple[bytes, str]:
    """在线程池内读取已生成的 Word 文件及原始文件名。"""
    with open(result_path, "rb") as result_file:
        filled_bytes = result_file.read()

    # 下载文件名需要查询原始文档，使用独立会话避免跨线程复用请求会话。
    from app.db.session import SessionLocal

    read_db = SessionLocal()
    try:
        _, raw_filename, _ = bid_format_extractor_service.extract_and_export_bid_format(
            db=read_db,
            doc_id=document_id,
            user_id=user_id,
            tenant_id=tenant_id,
        )
    finally:
        read_db.close()

    return filled_bytes, f"【ReActAgent智能填报】{raw_filename or '投标文件格式.docx'}"


@router.get("/agent-fill-bid-format/{document_id}/download")
async def download_agent_filled_bid_format(
    document_id: str,
    db: Session = Depends(deps.get_db),
    current_user: Optional[User] = Depends(deps.get_current_user_optional)
):
    """
    下载 BidFillerAgent (LangGraph + ReAct Agent) 自动填报完成后的 Word (.docx) 文档。
    """
    if not document_id:
        raise HTTPException(status_code=400, detail="未提供有效的 document_id 参数")

    logger.info(f"🤖 收到 BidFillerAgent 填报 Word 下载请求: doc_id={document_id}")

    drafts_dir = os.path.join(os.getcwd(), "uploads", "drafts")
    result_path = os.path.join(drafts_dir, f"agent_fill_result_{document_id[:8]}.docx")
    if not _has_non_empty_result_file(result_path):
        logger.warning(f"拦截未完成标书下载，不自动启动撰写任务: document_id={document_id}")
        raise HTTPException(
            status_code=409,
            detail="尚未生成已填写标书，请先填写标书后再下载。",
        )

    logger.info(f"   📄 命中已生成的结果文件: {result_path}")
    try:
        filled_bytes, out_filename = await run_in_threadpool(
            _read_agent_result_file,
            result_path,
            document_id,
            current_user.id if hasattr(current_user, "id") else None,
            current_user.tenant_id if hasattr(current_user, "tenant_id") else None,
        )
    except Exception as download_error:
        logger.exception(f"读取 ReActAgent 结果文件失败: {download_error}")
        raise HTTPException(status_code=500, detail="读取已生成的标书 Word 文件失败") from download_error

    return Response(
        content=filled_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(out_filename)}",
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )


# ============================================================
# 2. 单层通用参数 Endpoint (必须位于多层子路径之后)
# ============================================================

@router.get("/extract-bid-format/{document_id}")
@router.post("/extract-bid-format/{document_id}")
async def extract_and_download_bid_format(
    document_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user)
):
    """
    定位提取招投标原始文档中的“投标文件格式/响应格式”全量内容，并自动导出为 Word (.docx) 文件供下载。
    """
    if not document_id:
        raise HTTPException(status_code=400, detail="未提供有效的 document_id 参数")

    logger.info(f"收到投标文件格式提取请求: doc_id={document_id}, user_id={current_user.id}")

    try:
        docx_bytes, filename, mode = await run_in_threadpool(
            bid_format_extractor_service.extract_and_export_bid_format,
            db=db,
            doc_id=document_id,
            user_id=current_user.id if hasattr(current_user, 'id') else None,
            tenant_id=current_user.tenant_id if hasattr(current_user, 'tenant_id') else None,
            force_reextract=True,
        )

        if not docx_bytes:
            raise HTTPException(status_code=500, detail="未提取到有效内容或 Word 生成失败")

        encoded_filename = urllib.parse.quote(filename)

        headers = {
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
            "X-Extraction-Mode": mode,
            "Access-Control-Expose-Headers": "Content-Disposition, X-Extraction-Mode"
        }

        return Response(
            content=docx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers=headers
        )

    except HTTPException:
        raise
    except ModelUnavailableError as model_error:
        logger.warning("投标文件格式提取所需模型不可用: doc_id={}, error={}", document_id, model_error)
        raise HTTPException(status_code=503, detail=str(model_error)) from model_error
    except FileNotFoundError as e:
        logger.warning(f"提取文件未找到: {str(e)}")
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        logger.exception(f"提取并导出投标文件格式发生未预期的异常: {str(e)}")
        raise HTTPException(status_code=500, detail=f"提取生成失败: {str(e)}")


@router.get("/download-bid-format-template/{document_id}")
async def download_bid_format_template(
    document_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """下载原格式投标文件模板；没有缓存时先执行一次提取。"""
    if not document_id:
        raise HTTPException(status_code=400, detail="未提供有效的 document_id 参数")

    logger.info("收到投标文件模板下载请求: doc_id={}, user_id={}", document_id, current_user.id)
    try:
        docx_bytes, filename, mode = await run_in_threadpool(
            bid_format_extractor_service.extract_and_export_bid_format,
            db=db,
            doc_id=document_id,
            user_id=current_user.id if hasattr(current_user, "id") else None,
            tenant_id=current_user.tenant_id if hasattr(current_user, "tenant_id") else None,
            force_reextract=False,
        )
        if not docx_bytes:
            raise HTTPException(status_code=500, detail="未提取到有效的投标文件模板")

        encoded_filename = urllib.parse.quote(filename)
        return Response(
            content=docx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
                "X-Extraction-Mode": mode,
                "Access-Control-Expose-Headers": "Content-Disposition, X-Extraction-Mode",
            },
        )
    except HTTPException:
        raise
    except ModelUnavailableError as model_error:
        logger.warning("投标文件模板下载所需模型不可用: doc_id={}, error={}", document_id, model_error)
        raise HTTPException(status_code=503, detail=str(model_error)) from model_error
    except FileNotFoundError as file_error:
        logger.warning("投标文件模板下载原文件未找到: doc_id={}, error={}", document_id, file_error)
        raise HTTPException(status_code=404, detail=str(file_error)) from file_error
    except Exception as download_error:
        logger.exception("下载投标文件模板发生未预期异常: doc_id={}, error={}", document_id, download_error)
        raise HTTPException(status_code=500, detail=f"投标文件模板下载失败: {download_error}") from download_error


@router.post("/reextract-bid-format/{document_id}")
async def reextract_bid_format_template(
    document_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """重新提取并刷新投标文件模板缓存，但不直接触发文件下载。"""
    if not document_id:
        raise HTTPException(status_code=400, detail="未提供有效的 document_id 参数")

    logger.info("收到投标文件模板重新提取请求: doc_id={}, user_id={}", document_id, current_user.id)
    try:
        template_bytes, filename, mode = await run_in_threadpool(
            bid_format_extractor_service.extract_and_export_bid_format,
            db=db,
            doc_id=document_id,
            user_id=current_user.id if hasattr(current_user, "id") else None,
            tenant_id=current_user.tenant_id if hasattr(current_user, "tenant_id") else None,
            force_reextract=True,
        )
        if not template_bytes:
            raise HTTPException(status_code=500, detail="未生成有效的投标文件模板")

        return success_response(
            data={"filename": filename, "extraction_mode": mode},
            message="投标文件模板重新提取成功",
        )
    except HTTPException:
        raise
    except ModelUnavailableError as model_error:
        logger.warning("投标文件模板重新提取所需模型不可用: doc_id={}, error={}", document_id, model_error)
        raise HTTPException(status_code=503, detail=str(model_error)) from model_error
    except FileNotFoundError as file_error:
        logger.warning("投标文件模板重新提取原文件未找到: doc_id={}, error={}", document_id, file_error)
        raise HTTPException(status_code=404, detail=str(file_error)) from file_error
    except Exception as extract_error:
        logger.exception("重新提取投标文件模板发生未预期异常: doc_id={}, error={}", document_id, extract_error)
        raise HTTPException(status_code=500, detail=f"投标文件模板重新提取失败: {extract_error}") from extract_error


@router.get("/fill-bid-format/{document_id}")
@router.post("/fill-bid-format/{document_id}")
async def fill_and_download_bid_format(
    document_id: str,
    request_data: Optional[dict] = None,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user)
):
    """
    提取并导出 100% 原汁原味的《投标文件格式》Word (.docx) 文档。
    """
    if not document_id:
        raise HTTPException(status_code=400, detail="未提供有效的 document_id 参数")

    logger.info(f"收到投标文件格式纯净导出请求: doc_id={document_id}")

    try:
        template_bytes, filename, mode = await run_in_threadpool(
            bid_format_extractor_service.extract_and_export_bid_format,
            db=db,
            doc_id=document_id,
            user_id=current_user.id if hasattr(current_user, 'id') else None,
            tenant_id=current_user.tenant_id if hasattr(current_user, 'tenant_id') else None
        )

        if not template_bytes:
            raise HTTPException(status_code=500, detail="未提取到原始格式模版")

        output_filename = f"【投标文件格式】{filename}"
        encoded_filename = urllib.parse.quote(output_filename)

        headers = {
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
            "X-Extraction-Mode": mode,
            "Access-Control-Expose-Headers": "Content-Disposition, X-Extraction-Mode"
        }

        return Response(
            content=template_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers=headers
        )

    except HTTPException:
        raise
    except ModelUnavailableError as model_error:
        logger.warning("投标文件格式导出所需模型不可用: doc_id={}, error={}", document_id, model_error)
        raise HTTPException(status_code=503, detail=str(model_error)) from model_error
    except Exception as e:
        logger.exception(f"导出投标文件格式发生未预期异常: {str(e)}")
        raise HTTPException(status_code=500, detail=f"投标文件格式导出失败: {str(e)}")


def _run_agent_bid_filling_in_background(
    document_id: str,
    u_id: str,
    t_id: str,
    custom_instructions: Optional[str] = None,
    category_hints: Optional[dict] = None,
    profile_id: Optional[str] = None,
    template_id: Optional[str] = None,
):
    """后台工作线程：执行长耗时的 BidFillerAgent 多 Agent 标书撰写与落盘"""
    from app.core.context import current_user_id, current_tenant_id, current_task_id
    from app.db.session import SessionLocal
    from app.schemas.bid_filler_schema import CompanyProfile
    from app.agents.bid_filler_agent import bid_filler_agent
    from app.services.bid_format_extractor_service import bid_format_extractor_service

    token_task = current_task_id.set(document_id)
    token_u = current_user_id.set(u_id)
    token_t = current_tenant_id.set(t_id)
    # 设置 Agent 工具查询使用的企业档案
    from app.agents.tools.bid_db_tools import current_profile_id as ctx_profile_id
    token_profile = ctx_profile_id.set(profile_id) if profile_id else None
    db: Session = SessionLocal()
    import time as _bg_time
    bg_start_t = _bg_time.time()
    try:
        template_bytes, filename, _ = bid_format_extractor_service.extract_and_export_bid_format(
            db=db, doc_id=document_id, user_id=None, tenant_id=None, template_id=template_id
        )
        if not template_bytes:
            logger.error(f"后台任务提取《投标文件格式》模板失败: doc_id={document_id}")
            from app.db.models.audit import AgentAuditLog
            err_log = AgentAuditLog(
                task_id=document_id,
                tenant_id=t_id,
                user_id=u_id,
                node_name="Supervisor-总控调度",
                action_type="llm_call_supervisor",
                status="failed",
                inputs={"chapter_title": "Supervisor-总控调度"},
                outputs={"summary": "❌ 后台提取《投标文件格式》模板失败"}
            )
            db.add(err_log)
            db.commit()
            return

        replacement_map, audit_report, filled_bytes = bid_filler_agent.process_filling_tasks(
            db=db,
            document_id=document_id,
            profile=CompanyProfile(),
            detected_placeholders=[],
            original_docx=template_bytes,
            custom_instructions=custom_instructions,
            category_hints=category_hints,
            profile_id=profile_id,
        )

        if not filled_bytes:
            from app.services.bid_format_filler_service import bid_format_filler_service as filler_svc
            filled_bytes = filler_svc.fill_docx_with_audit_trail(
                docx_bytes=template_bytes,
                replacement_map=replacement_map,
                audit_items=audit_report.audit_items if audit_report else []
            )

        if filled_bytes:
            drafts_dir = os.path.join(os.getcwd(), "uploads", "drafts")
            os.makedirs(drafts_dir, exist_ok=True)
            result_path = os.path.join(drafts_dir, f"agent_fill_result_{document_id[:8]}.docx")
            draft_path = os.path.join(drafts_dir, f"draft_{document_id}.docx")
            for p in [result_path, draft_path]:
                with open(p, "wb") as f:
                    f.write(filled_bytes)
            logger.info(f"✅ 后台标书撰写完成并已保存至: {result_path}")

            # 写入 Supervisor 最终完成日志，记录端到端真实物理总耗时
            total_wall_ms = int((_bg_time.time() - bg_start_t) * 1000)
            _persist_first_bid_fill_duration(db, document_id, total_wall_ms)
            try:
                from app.db.models.audit import AgentAuditLog
                manual_chapters_payload = [
                    item.model_dump() if hasattr(item, "model_dump") else item.dict()
                    for item in (audit_report.manual_chapters if audit_report else [])
                ]
                final_sup_log = AgentAuditLog(
                    task_id=document_id,
                    tenant_id=t_id,
                    user_id=u_id,
                    node_name="Supervisor-总控调度",
                    action_type="llm_call_supervisor",
                    status="master_completed",
                    execution_time_ms=total_wall_ms,
                    inputs={"document_id": document_id, "chapter_title": "Supervisor-总控调度", "wall_time_ms": total_wall_ms},
                    outputs={
                        "summary": f"✨ AI 团队自主撰写与原位写盘已全量收官！全流程耗时 {total_wall_ms / 1000:.1f} 秒。所有章节卡片均已更新。",
                        "manual_chapters": manual_chapters_payload,
                        "manual_pending_count": len(manual_chapters_payload),
                    }
                )
                db.add(final_sup_log)
                db.commit()
            except Exception as final_log_err:
                logger.warning(f"写入最终 Supervisor 完结日志异常: {final_log_err}")
    except Exception as e:
        logger.exception(f"❌ 后台标书撰写任务异常: {e}")
        try:
            from app.db.models.audit import AgentAuditLog
            err_log = AgentAuditLog(
                task_id=document_id,
                tenant_id=t_id,
                user_id=u_id,
                node_name="Supervisor-总控调度",
                action_type="llm_call_supervisor",
                status="failed",
                inputs={"chapter_title": "Supervisor-总控调度"},
                outputs={"summary": f"❌ 后台标书撰写任务异常中断: {str(e)}"}
            )
            db.add(err_log)
            db.commit()
        except Exception:
            pass
    finally:
        try:
            current_task_id.reset(token_task)
            current_user_id.reset(token_u)
            current_tenant_id.reset(token_t)
            if token_profile is not None:
                ctx_profile_id.reset(token_profile)
        except Exception:
            pass
        db.close()


def _dispatch_agent_bid_filling(
    document_id: str,
    user_id: str,
    tenant_id: str,
    custom_instructions: Optional[str],
    category_hints: Optional[dict],
    profile_id: Optional[str],
    template_id: Optional[str],
    db: Session,
) -> tuple[Optional[int], str]:
    """在线程池内完成标书撰写任务的同步调度，避免阻塞 FastAPI 事件循环。"""
    from app.services.bid_fill_task_service import bid_fill_task_service, start_bid_fill_process

    reservation, reservation_status = bid_fill_task_service.acquire(document_id)
    if reservation is None:
        return None, reservation_status

    try:
        # 清理该文档上一次的填报审计日志，并立即注入全局起始 in_progress 记录。
        try:
            from app.db.models.audit import AgentAuditLog
            from sqlalchemy import cast, String

            db.query(AgentAuditLog).filter(
                or_(
                    AgentAuditLog.task_id == document_id,
                    cast(AgentAuditLog.inputs, String).like(f"%{document_id}%"),
                )
            ).delete(synchronize_session=False)

            init_log = AgentAuditLog(
                task_id=document_id,
                tenant_id=tenant_id,
                user_id=user_id,
                node_name="Supervisor-总控调度",
                action_type="llm_call_supervisor",
                status="in_progress",
                inputs={
                    "document_id": document_id,
                    "chapter_title": "Supervisor-总控调度",
                    "msg": "准备启动新一轮 Agent 全自主起草...",
                },
                outputs={"summary": "正在初始化 Agent 专家团队与指令解析..."},
            )
            db.add(init_log)
            db.commit()
            logger.info(f"成功清理旧日志并写入初始 in_progress 状态记录，doc_id={document_id}")
        except Exception as cleanup_error:
            # 历史审计日志清理失败不应阻止新任务启动，但必须回滚并记录原因。
            logger.warning(f"清理旧 AuditLog 异常: {cleanup_error}")
            db.rollback()

        process_id = start_bid_fill_process(
            document_id=document_id,
            user_id=user_id,
            tenant_id=tenant_id,
            custom_instructions=custom_instructions,
            category_hints=category_hints,
            reservation_data=reservation.to_payload(),
            profile_id=profile_id,
            template_id=template_id,
        )
        return process_id, reservation_status
    except Exception:
        # 子进程启动失败时释放已预留槽位，避免后续任务永久收到容量已满。
        bid_fill_task_service.release(reservation)
        raise


@router.post("/agent-fill-bid-format/{document_id}")
async def trigger_agent_bid_filling(
    document_id: str,
    request_body: Optional[BidFillRequest] = None,
    db: Session = Depends(deps.get_db),
    current_user: Optional[User] = Depends(deps.get_current_user_optional)
):
    """
    触发 BidFillerAgent (LangGraph + ReAct Agent) 自动填报。
    使用独立子进程隔离长耗时 Word 与 Agent 操作，配合 SSE 获得实时进度。
    """
    if not document_id:
        raise HTTPException(status_code=400, detail="未提供有效的 document_id 参数")

    logger.info(
        "🤖 收到 BidFillerAgent ReAct 自动填报请求: doc_id={}, profile_id={}",
        document_id,
        request_body.profile_id if request_body else None,
    )

    u_id = current_user.id if (current_user and hasattr(current_user, 'id')) else "default-user"
    t_id = current_user.tenant_id if (current_user and hasattr(current_user, 'tenant_id')) else "default-tenant"

    custom_instructions = None
    category_hints = None
    profile_id = None
    template_id = None
    if request_body:
        custom_instructions = request_body.custom_instructions
        category_hints = request_body.category_hints
        profile_id = request_body.profile_id
        template_id = request_body.template_id

    try:
        # Redis、数据库和 Windows 子进程启动都是同步操作，统一移出事件循环。
        process_id, reservation_status = await run_in_threadpool(
            _dispatch_agent_bid_filling,
            document_id,
            u_id,
            t_id,
            custom_instructions,
            category_hints,
            profile_id,
            template_id,
            db,
        )
    except Exception as dispatch_error:
        logger.exception(f"启动独立标书撰写进程失败: document_id={document_id}, error={dispatch_error}")
        raise HTTPException(status_code=503, detail="标书撰写进程启动失败，请稍后重试") from dispatch_error

    if process_id is None:
        reservation_messages = {
            "document_running": "该标书正在撰写中，请勿重复提交",
            "capacity_reached": "当前已有标书撰写任务正在执行，请稍后重试",
            "redis_unavailable": "任务调度服务暂不可用，请检查 Redis 后重试",
        }
        status_code = 503 if reservation_status == "redis_unavailable" else 409
        raise HTTPException(
            status_code=status_code,
            detail=reservation_messages[reservation_status],
        )

    return {
        "document_id": document_id,
        "task_id": f"process-{process_id}",
        "process_id": process_id,
        "status": "processing",
        "message": "已成功启动独立进程执行 Agent 团队标书撰写，请通过 SSE 实时监听进度"
    }
