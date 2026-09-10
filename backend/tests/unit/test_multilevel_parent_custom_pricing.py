"""
单元测试：多级 BOM 成套统价锁定与总额防双重计费测试 (test_multilevel_parent_custom_pricing.py)

验证多级 BOM 场景下（如 L1 -> L2 -> L3）：
1. 母项 L1 设置成套自定义统价时，其下属中间成套节点 L2 与底层节点 L3 均被统价锁定；
2. rollup_hierarchical_cost_items 层级汇总服务计算的总成本严格为 L1 统价，杜绝双重计费；
3. bom_export_service 导出 Word/Excel 时，被锁定的中间项与子项单价和小计规范渲染为 '--'，备注标记'已统入成套价'，总价不发生累加；
4. 母项未设统价而中间总成设统价时，锁定状态与汇总正确向下传递。
"""

import io
import pytest
from docx import Document
import openpyxl

from app.services.cost_service import rollup_hierarchical_cost_items
from app.services.bom_export_service import (
    _calculate_export_total,
    _collect_export_rows,
    generate_bom_docx,
    generate_bom_xlsx,
)


def _build_three_level_fixture():
    """构造标准三级 BOM 数据：L1 电气二次部分 -> L2 变电站自动化系统 -> L3 监控终端。"""
    return [
        {
            "node_id": "node-l1",
            "name": "电气二次部分",
            "qty": 1.0,
            "unit": "套",
            "ref_price": 10.0,
            "subtotal": 10.0,
            "pricing_mode": "parent",
            "is_parent_modified": True,
            "match_quality": "手动修改",
            "parent_node_id": None,
            "parent_item": None,
        },
        {
            "node_id": "node-l2",
            "name": "变电站自动化系统",
            "qty": 1.0,
            "unit": "套",
            "ref_price": 7500.0,
            "subtotal": 7500.0,
            "pricing_mode": "parent",
            "is_parent_modified": True,
            "match_quality": "手动修改",
            "parent_node_id": "node-l1",
            "parent_item": "电气二次部分",
        },
        {
            "node_id": "node-l3",
            "name": "站控层监控主机",
            "qty": 2.0,
            "unit": "台",
            "ref_price": 100.0,
            "subtotal": 200.0,
            "pricing_mode": "item",
            "is_parent_modified": False,
            "match_quality": "精准匹配",
            "parent_node_id": "node-l2",
            "parent_item": "变电站自动化系统",
        },
    ]


def test_multilevel_parent_pricing_rollup_and_lock():
    """测试多级统价层级汇总：L1 统价 10 元时，L2 与 L3 被锁定且总额严格为 10 元。"""
    items = _build_three_level_fixture()
    clean_items, total_cost, unmatched_count = rollup_hierarchical_cost_items(items)

    assert total_cost == 10.0, f"期望预估总成本为 10.00，实际为 {total_cost}"
    assert unmatched_count == 0

    item_by_id = {item["node_id"]: item for item in clean_items}

    # L1 为顶层统价母项
    assert item_by_id["node-l1"]["is_locked_by_parent"] is False
    assert item_by_id["node-l1"]["subtotal"] == 10.0

    # L2 为被 L1 锁定的中间总成
    assert item_by_id["node-l2"]["is_locked_by_parent"] is True
    assert item_by_id["node-l2"]["subtotal"] == 0.0

    # L3 为被 L1 间接锁定的底层元件
    assert item_by_id["node-l3"]["is_locked_by_parent"] is True
    assert item_by_id["node-l3"]["subtotal"] == 0.0


def test_multilevel_parent_pricing_bom_export_total_no_double_count():
    """测试导出服务展开导出行时，总额严格杜绝 10 + 7500 = 7510 的双重计费。"""
    items = _build_three_level_fixture()
    # 先跑一次 rollup 打上锁定标记与树关系
    clean_items, _, _ = rollup_hierarchical_cost_items(items)
    export_rows = _collect_export_rows(clean_items)

    calculated_total = _calculate_export_total(export_rows)
    assert calculated_total == 10.0, f"导出总额计算错误，期望 10.0，实际 {calculated_total}"


def test_multilevel_parent_pricing_word_and_excel_export_locked_rendering():
    """测试导出 Word 和 Excel 时，被锁定中间项规范展示为 '--' 与备注标记。"""
    items = _build_three_level_fixture()
    clean_items, total_cost, _ = rollup_hierarchical_cost_items(items)

    # 1. 验证 Word 导出
    doc_io = generate_bom_docx(
        document_title="多级BOM统价测试标书",
        items=clean_items,
        total_cost=total_cost,
    )
    doc = Document(io.BytesIO(doc_io.getvalue()))
    assert len(doc.tables) >= 1
    table = doc.tables[0]

    # 表头(0) + L1(1) + L2(2) + L3(3) + 表尾(4)
    assert len(table.rows) == 5

    # L1 行数据：单价 10.00，总价 10.00
    l1_row_cells = [c.text for c in table.rows[1].cells]
    assert "电气二次部分" in l1_row_cells[1]
    assert "10.00" in l1_row_cells[6]
    assert "10.00" in l1_row_cells[7]

    # L2 行数据：被母项锁定，单价为 --，总价为 --，备注含'已统入成套价'
    l2_row_cells = [c.text for c in table.rows[2].cells]
    assert "变电站自动化系统" in l2_row_cells[1]
    assert l2_row_cells[6] == "--"
    assert l2_row_cells[7] == "--"
    assert "已统入成套价" in l2_row_cells[8]

    # L3 行数据：被母项间接锁定，单价为 --，总价为 --，备注含'已统入成套价'
    l3_row_cells = [c.text for c in table.rows[3].cells]
    assert "站控层监控主机" in l3_row_cells[1]
    assert l3_row_cells[6] == "--"
    assert l3_row_cells[7] == "--"
    assert "已统入成套价" in l3_row_cells[8]

    # 表尾总价为 10.00 与大写拾元整
    footer_row_cells = [c.text for c in table.rows[4].cells]
    assert "10.00" in footer_row_cells[7]
    assert "拾元整" in footer_row_cells[0]

    # 2. 验证 Excel 导出
    xlsx_io = generate_bom_xlsx(
        document_title="多级BOM统价测试标书",
        items=clean_items,
        total_cost=total_cost,
    )
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_io.getvalue()))
    ws = wb.active

    # Row 6: L1, Row 7: L2, Row 8: L3, Row 9: Footer
    # Col 7: 单价, Col 8: 总价, Col 9: 备注
    assert ws.cell(row=6, column=2).value == "电气二次部分"
    assert ws.cell(row=6, column=7).value == 10.0
    assert ws.cell(row=6, column=8).value == 10.0

    assert ws.cell(row=7, column=2).value == "变电站自动化系统"
    assert ws.cell(row=7, column=7).value == "--"
    assert ws.cell(row=7, column=8).value == "--"
    assert "已统入成套价" in str(ws.cell(row=7, column=9).value)

    assert ws.cell(row=8, column=2).value == "站控层监控主机"
    assert ws.cell(row=8, column=7).value == "--"
    assert ws.cell(row=8, column=8).value == "--"
    assert "已统入成套价" in str(ws.cell(row=8, column=9).value)

    assert ws.cell(row=9, column=8).value == 10.0


def test_multilevel_parent_pricing_fallback_when_root_is_sum_mode():
    """测试当 L1 为普通汇总模式时，L2 的统价 7500 元生效并向上汇总为 L1 的 7500 元。"""
    items = [
        {
            "node_id": "node-l1",
            "name": "电气二次部分",
            "qty": 1.0,
            "unit": "套",
            "ref_price": 0.0,
            "subtotal": 0.0,
            "pricing_mode": "item",  # 普通汇总模式
            "is_parent_modified": False,
            "parent_node_id": None,
            "parent_item": None,
        },
        {
            "node_id": "node-l2",
            "name": "变电站自动化系统",
            "qty": 1.0,
            "unit": "套",
            "ref_price": 7500.0,
            "subtotal": 7500.0,
            "pricing_mode": "parent",
            "is_parent_modified": True,
            "match_quality": "手动修改",
            "parent_node_id": "node-l1",
            "parent_item": "电气二次部分",
        },
        {
            "node_id": "node-l3",
            "name": "站控层监控主机",
            "qty": 2.0,
            "unit": "台",
            "ref_price": 100.0,
            "subtotal": 200.0,
            "pricing_mode": "item",
            "is_parent_modified": False,
            "parent_node_id": "node-l2",
            "parent_item": "变电站自动化系统",
        },
    ]

    clean_items, total_cost, _ = rollup_hierarchical_cost_items(items)
    assert total_cost == 7500.0

    item_by_id = {item["node_id"]: item for item in clean_items}
    # L1 由 L2 汇总驱动为 7500
    assert item_by_id["node-l1"]["subtotal"] == 7500.0
    assert item_by_id["node-l1"]["is_locked_by_parent"] is False

    # L2 自身为成套统价，未被锁定
    assert item_by_id["node-l2"]["subtotal"] == 7500.0
    assert item_by_id["node-l2"]["is_locked_by_parent"] is False

    # L3 被 L2 锁定
    assert item_by_id["node-l3"]["subtotal"] == 0.0
    assert item_by_id["node-l3"]["is_locked_by_parent"] is True
